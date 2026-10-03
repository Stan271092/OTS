from pathlib import Path
from io import BytesIO
from collections import Counter
import sys, re, json, html
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, A3, A2, landscape
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Flowable, Paragraph, Table, TableStyle, Spacer, PageBreak, KeepTogether, NextPageTemplate
from svglib.svglib import svg2rlg
from reportlab.graphics import renderPDF
import pymupdf

sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path('D:/AI/ORCA PROJECTS/OTR/Сайт визитка/OTS-GitHub')
OUT=ROOT/'Клиентские каталоги 2026'
CHECK=Path('C:/Users/Stasius/AppData/Local/Temp/ots-catalog-whole-tables-review')
CHECK.mkdir(exist_ok=True)
SOURCE=ROOT/'Каталог продукции_2026 — v2.xlsx'
W,H=landscape(A4); M=18*mm; WIDTH=W-2*M
NAVY=colors.HexColor('#112543'); BLUE=colors.HexColor('#3A6CC0'); PALE=colors.HexColor('#F2F6FC'); LINE=colors.HexColor('#CDD9EB'); INK=colors.HexColor('#22334A'); MUTED=colors.HexColor('#61738D')
pdfmetrics.registerFont(TTFont('OTS','C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('OTS-B','C:/Windows/Fonts/arialbd.ttf'))
ST={
 'cell':ParagraphStyle('cell',fontName='OTS',fontSize=8.5,leading=11,textColor=INK,splitLongWords=True),
 'head':ParagraphStyle('head',fontName='OTS-B',fontSize=8,leading=10.2,textColor=colors.white),
 'h':ParagraphStyle('h',fontName='OTS-B',fontSize=17,leading=22,textColor=NAVY,spaceAfter=9,keepWithNext=True),
 'intro':ParagraphStyle('intro',fontName='OTS',fontSize=9.5,leading=13,textColor=MUTED,spaceAfter=10,keepWithNext=True),
 'note':ParagraphStyle('note',fontName='OTS',fontSize=8.5,leading=12,textColor=MUTED,spaceAfter=6),
}
ledger=[]

def text(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v).replace('\r','').strip()

def p(v,kind='cell'):
    return Paragraph(html.escape(text(v)).replace('\n','<br/>'),ST[kind])

def val(ws,r,c):
    cell=ws.cell(r,c)
    if cell.value is None:
        for region in ws.merged_cells.ranges:
            if region.min_row<=r<=region.max_row and region.min_col<=c<=region.max_col:
                cell=ws.cell(region.min_row,region.min_col);break
    if cell.value is not None:ledger.append((ws.title,cell.coordinate,text(cell.value)))
    return text(cell.value)

def table(headers,rows,weights,spans=(),available=WIDTH):
    widths=[available*x/sum(weights) for x in weights]
    assert abs(sum(widths)-available)<0.01
    assert all(len(r)==len(headers)==len(widths) for r in rows)
    t=Table([[p(v,'head') for v in headers]]+[[p(v) for v in row] for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([
      ('BACKGROUND',(0,0),(-1,0),NAVY),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,PALE]),
      ('LINEBELOW',(0,0),(-1,0),.7,BLUE),('LINEBELOW',(0,1),(-1,-1),.35,LINE),
      ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
      ('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]+list(spans)))
    return t

class RotatedHeader(Flowable):
    def __init__(self,value,font_size=7.3):
        super().__init__()
        self.lines=[line.strip() for line in value.split('\n')]
        self.font_size=font_size
        self.leading=font_size+1.7
        self.width=len(self.lines)*self.leading+2
        self.height=max(pdfmetrics.stringWidth(line,'OTS-B',font_size) for line in self.lines)+2
        self.hAlign='CENTER'
    def draw(self):
        c=self.canv
        c.saveState()
        c.translate(self.width,0)
        c.rotate(90)
        c.setFont('OTS-B',self.font_size)
        c.setFillColor(colors.white)
        for i,line in enumerate(self.lines):
            c.drawString(1,self.width-self.font_size-1-i*self.leading,line)
        c.restoreState()

COMPOUND_PANELS=(tuple(range(2,20)),tuple(range(20,31)))
COMPOUND_WEIGHTS=[39,28,38,42]+[14]*10+[25,25,25,24,25,22,22,22,23,16,25,22,24,25,56]

SPREAD_GUTTER=16*mm
PANEL_WIDTH=(landscape(A2)[0]-2*M-SPREAD_GUTTER)/2

def compound_table(ws,cols,rows,row_heights=None):
    # The two panels share the same header and body row heights, as in a printed spread.
    source_rows=[2,3,4,5]+list(rows)
    row_index={r:i for i,r in enumerate(source_rows)}
    col_index={c:i for i,c in enumerate(cols)}
    weights=[COMPOUND_WEIGHTS[c-2] for c in cols]
    available=PANEL_WIDTH
    widths=[available*v/sum(weights) for v in weights]
    headstyle=ParagraphStyle('wide-head',parent=ST['head'],fontSize=7.3,leading=9)
    cellstyle=ParagraphStyle('wide-cell',parent=ST['cell'],fontSize=7.7,leading=10)
    cells=[]
    for r in source_rows:
        row=[]
        for c in cols:
            value=val(ws,r,c) if ws.cell(r,c).value is not None else ''
            if r==3 and 6<=c<=15 and ws.cell(r,c).alignment.textRotation==90:
                row.append(RotatedHeader(value))
            else:
                row.append(Paragraph(html.escape(value).replace('\n','<br/>'),headstyle if r<=5 else cellstyle))
        cells.append(row)
    t=Table(cells,colWidths=widths,rowHeights=row_heights,repeatRows=4,hAlign='LEFT')
    commands=[('BACKGROUND',(0,0),(-1,3),NAVY),
      ('GRID',(0,0),(-1,-1),.3,LINE),('VALIGN',(0,0),(-1,-1),'TOP'),
      ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),
      ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]
    for i,r in enumerate(rows,4):
        commands.append(('BACKGROUND',(0,i),(-1,i),colors.white if (r-7)%2==0 else PALE))
    for c in cols:
        if 6<=c<=15:
            commands.append(('VALIGN',(col_index[c],1),(col_index[c],3),'MIDDLE'))
    for region in ws.merged_cells.ranges:
        if region.min_row>=2 and region.max_row<=5:
            visible=[c for c in cols if region.min_col<=c<=region.max_col]
            if visible:
                commands.append(('SPAN',(col_index[visible[0]],row_index[region.min_row]),(col_index[visible[-1]],row_index[region.max_row])))
    t.setStyle(TableStyle(commands))
    return t

class CompoundSpread(Flowable):
    def __init__(self,ws,rows,row_heights):
        super().__init__()
        self.width=2*PANEL_WIDTH+SPREAD_GUTTER
        self.height=sum(row_heights)
        self.panels=[compound_table(ws,cols,rows,row_heights) for cols in COMPOUND_PANELS]
    def draw(self):
        for side,panel in enumerate(self.panels):
            width,height=panel.wrap(PANEL_WIDTH,self.height)
            assert abs(height-self.height)<.01
            panel.drawOn(self.canv,side*(PANEL_WIDTH+SPREAD_GUTTER),0)

def compound_spread(ws,rows,row_heights):
    return CompoundSpread(ws,rows,row_heights)

def compound_pages(ws,title,ru):
    rows=list(range(7,40))
    measured=[]
    for cols in COMPOUND_PANELS:
        t=compound_table(ws,cols,rows)
        t.wrap(PANEL_WIDTH,100000)
        measured.append(t._rowHeights)
    heights=[max(pair) for pair in zip(*measured)]
    available=landscape(A2)[1]-36*mm
    # Match the reference spreads: ARC-15 through C-59, then ARL-111 through C-51.
    chunks=[(0,16),(16,len(rows))]
    for start,end in chunks:
        assert sum(heights[:4])+sum(heights[4+start:4+end])<=available,('Reference spread does not fit',ws.title,start,end)
    story=[]
    for part,(start,end) in enumerate(chunks):
        if story:story.append(PageBreak())
        story.append(compound_spread(ws,rows[start:end],heights[:4]+heights[4+start:4+end]))
    return story

class Cover(Flowable):
    width=WIDTH; height=465
    def __init__(self,ru):super().__init__();self.ru=ru;self.width=WIDTH;self.height=465
    def draw(self):
        c=self.canv
        d=svg2rlg(str(ROOT/('new logo RU.svg' if self.ru else 'new logo eng.svg')))
        scale=185/d.width
        c.saveState();c.translate(0,350);c.scale(scale,scale);renderPDF.draw(d,c,0,0);c.restoreState()
        c.setStrokeColor(BLUE);c.setLineWidth(3);c.line(0,290,56,290)
        c.setFillColor(NAVY);c.setFont('OTS-B',34)
        c.drawString(0,228,'Каталог продукции' if self.ru else 'Product Catalogue')
        c.setFont('OTS',12);c.setFillColor(MUTED)
        for i,s in enumerate(['Компаунды','Эпоксидные смолы','Отвердители'] if self.ru else ['Compounds','Epoxy resins','Hardeners']):c.drawString(0,116-i*22,s)

class Divider(Flowable):
    width=WIDTH;height=465
    def __init__(self,title,sub,number):super().__init__();self.title=title;self.sub=sub;self.number=number;self.width=WIDTH;self.height=465
    def draw(self):
        c=self.canv;c.setFillColor(NAVY);c.rect(0,45,WIDTH,365,fill=1,stroke=0)
        c.setFillColor(BLUE);c.rect(0,45,8,365,fill=1,stroke=0)
        c.setFont('OTS',12);c.setFillColor(colors.HexColor('#AFC9EF'));c.drawString(30,363,self.number+' / 03')
        c.setFont('OTS-B',29);c.setFillColor(colors.white);c.drawString(30,237,self.title)
        c.setFont('OTS',11);c.setFillColor(colors.HexColor('#DCE8FA'));c.drawString(30,198,self.sub)

class Catalog(BaseDocTemplate):
    def __init__(self,path,title,**kwargs):
        super().__init__(str(path),pagesize=(W,H),title=title,author='OTS',**kwargs)
        self.section='';self.title_label=title;self.printed_page=1
        for name,size in [('main',(W,H)),('wide',landscape(A3)),('compound',landscape(A2))]:
            w,h=size
            self.addPageTemplates(PageTemplate(name,pagesize=size,frames=[Frame(M,20*mm,w-2*M,h-36*mm,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPageEnd=self.end_page))
    def afterFlowable(self,f):
        if isinstance(f,Cover):self.section=''
        if isinstance(f,Divider):
            self.section=f.title
            self.canv.bookmarkPage(f.number)
            self.canv.addOutlineEntry(f.title,f.number,0)
    def end_page(self,c,doc):
        width=c._pagesize[0]
        c.saveState();c.setStrokeColor(LINE);c.setLineWidth(.5);c.line(M,13*mm,width-M,13*mm)
        c.setFont('OTS',8);c.setFillColor(MUTED)
        c.drawString(M,8*mm,self.title_label+('  /  '+self.section if self.section else ''))
        if doc.pageTemplate.id=='compound':
            c.drawRightString(M+PANEL_WIDTH,8*mm,str(self.printed_page))
            c.drawRightString(width-M,8*mm,str(self.printed_page+1))
            self.printed_page+=2
        else:
            c.drawRightString(width-M,8*mm,str(self.printed_page))
            self.printed_page+=1
        c.restoreState()

def heading(s,sub=None):return [p(s,'h')]+([p(sub,'intro')] if sub else [])

def header(ws,c):
    return '\n'.join(filter(None,[val(ws,r,c) for r in (3,4,5) if ws.cell(r,c).value is not None]))

def build(lang):
    global ledger
    ledger=[];ru=lang=='RU'
    wb=load_workbook(SOURCE,data_only=False)
    title='Каталог продукции' if ru else 'Product Catalogue'
    dest=OUT/('Каталог_продукции.pdf' if ru else 'Product_Catalogue.pdf')
    story=[Cover(ru),PageBreak()]
    comp=wb['Компаунды RU' if ru else 'Compounds ENG']
    sections=['Компаунды','Эпоксидные смолы','Отвердители'] if ru else ['Compounds','Epoxy resins','Hardeners']
    story.extend([Divider(sections[0],'Системы, переработка и технические свойства' if ru else 'Systems, processing and technical properties','01'),NextPageTemplate('compound'),PageBreak()])
    story+=compound_pages(comp,sections[0],ru)
    extra=[[val(comp,r,2),val(comp,r,3)] for r in range(46,49) if comp.cell(r,2).value is not None]
    if extra:
        story+=[Spacer(1,6)]
        story.extend(p(label+': '+description,'note') for label,description in extra)
    story.append(NextPageTemplate('main'))
    story+=[PageBreak(),Divider(sections[1],'Базовые и специальные эпоксидные смолы' if ru else 'Base and specialty epoxy resins','02'),NextPageTemplate('wide'),PageBreak()]
    ws=wb['Эпоксидные смолы RU' if ru else 'Epoxy resins EN']
    for i,(tr,dr,hr,rs,cs,weights) in enumerate([
       (1,2,3,range(4,12),range(1,8),[13,25,13,9,10,13,24]),
       (13,14,15,range(16,22),range(1,9),[13,18,10,9,10,13,10,23]),
       (23,24,25,range(26,40),range(1,9),[12,22,12,9,10,15,10,24]),
    ]):
        if i:story.append(PageBreak())
        story+=heading(val(ws,tr,1),val(ws,dr,1))
        rows=[[val(ws,r,c) if ws.cell(r,c).value is not None else '' for c in cs] for r in rs]
        spans=[]
        for region in ws.merged_cells.ranges:
            if region.min_row in rs and region.max_row in rs:
                spans.append(('SPAN',(region.min_col-1,region.min_row-rs.start+1),(region.max_col-1,region.max_row-rs.start+1)))
        story.append(table([val(ws,hr,c) for c in cs],rows,weights,spans,available=landscape(A3)[0]-2*M))
    story+=[NextPageTemplate('main'),PageBreak(),Divider(sections[2],'Технические характеристики отвердителей' if ru else 'Technical data of curing agents','03'),PageBreak()]
    ws=wb['Отвердители RU' if ru else 'Hardeners EN']
    story+=heading(sections[2],'Технические характеристики отвердителей для различных режимов отверждения.' if ru else 'Technical data of curing agents for various curing regimes.')
    story.append(table([val(ws,2,c) for c in range(1,8)],[[val(ws,r,c) for c in range(1,8)] for r in range(3,8)],[26,13,24,12,11,10,13]))
    doc=Catalog(dest,title);doc.build(story)
    # Every data cell in every product row must have been included in the rendering input.
    used={(s,c) for s,c,v in ledger}
    for sn,rr,cc in [(comp.title,range(7,40),range(2,31)),(wb.worksheets[2 if ru else 3].title,list(range(4,12))+list(range(16,22))+list(range(26,40)),range(1,9)),(ws.title,range(3,8),range(1,8))]:
        for r in rr:
            for c in cc:
                cell=wb[sn].cell(r,c)
                if cell.value is not None:assert (sn,cell.coordinate) in used,('MISSING SOURCE',sn,cell.coordinate)
    pdf=pymupdf.open(dest)
    normalized=lambda s:re.sub(r'\s+','',s)
    extracted=normalized(''.join(pg.get_text() for pg in pdf))
    for sn,coord,value in ledger:
        assert normalized(value) in extracted,('MISSING PDF TEXT',lang,sn,coord,value)
    violations=[]
    for i,pg in enumerate(pdf):
        words=pg.get_text('words')
        assert len(words)>4,('BLANK',i+1)
        for x0,y0,x1,y1,*_ in words:
            if x0<M-1 or x1>pg.rect.width-M+1 or y0<10 or y1>pg.rect.height-12:violations.append((i+1,(x0,y0,x1,y1)))
    assert not violations,violations[:10]
    (CHECK/f'{lang}-ledger.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2),encoding='utf-8')
    for i,pg in enumerate(pdf):pg.get_pixmap(matrix=pymupdf.Matrix(min(1,1200/pg.rect.width),min(1,1200/pg.rect.width)),alpha=False).save(CHECK/f'{lang}-{i+1:02}.png')
    print('PASS',lang,'pages',len(pdf),'source entries',len(ledger),'all source data present in extracted PDF; page bounds OK')
    return len(pdf)

for lang in ('RU','EN'):build(lang)
